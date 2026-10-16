#!/usr/bin/env python3
"""
Exporteert alle openstaande issues uit SonarQube naar een Excel-bestand (.xlsx).
Slaat het bestand op in de map 'sonar_issues' met de datum in de bestandsnaam.
Verwijdert automatisch exports die ouder zijn dan 3 maanden.
"""

import os
import re
import sys
import json
import base64
import urllib.request
import urllib.error
from datetime import datetime, timedelta
from pathlib import Path

# Project root bepalen (één niveau boven 'scripts')
ROOT_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = ROOT_DIR / "sonar_issues"
ENV_FILE = ROOT_DIR / ".env"

DEFAULT_SONAR_URL = "http://192.168.1.250:9000"
DEFAULT_PROJECT_KEY = "Digidokters"
BEWAARTERMIJN_DAGEN = 90  # 3 maanden


def laad_env(env_pad: Path) -> dict:
    """Eenvoudige .env parser zonder externe dependencies."""
    env_vars = {}
    if not env_pad.exists():
        return env_vars

    with open(env_pad, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            if "=" in line:
                k, v = line.split("=", 1)
                k = k.strip()
                v = v.strip().strip("'\"")
                env_vars[k] = v
    return env_vars


def haal_alle_issues_op(sonar_url: str, project_key: str, token: str) -> list:
    """Haalt alle openstaande issues op via de SonarQube REST API met paginatie."""
    alle_issues = []
    page = 1
    page_size = 500

    auth_str = base64.b64encode(f"{token}:".encode()).decode()
    headers = {
        "Authorization": f"Basic {auth_str}",
        "Accept": "application/json"
    }

    print(f"Verbinden met SonarQube: {sonar_url} (Project: {project_key})...")

    while True:
        url = (
            f"{sonar_url.rstrip('/')}/api/issues/search?"
            f"componentKeys={project_key}&resolved=false&ps={page_size}&p={page}"
        )
        req = urllib.request.Request(url, headers=headers)

        try:
            with urllib.request.urlopen(req, timeout=30) as resp:
                data = json.loads(resp.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            raise RuntimeError(f"SonarQube API fout (HTTP {e.code}): {e.read().decode('utf-8')}") from e
        except Exception as e:
            raise RuntimeError(f"Fout bij verbinden met SonarQube: {e}") from e

        totaal = data.get("total", 0)
        issues_op_pagina = data.get("issues", [])
        alle_issues.extend(issues_op_pagina)

        print(f"  -> Pagina {page}: {len(issues_op_pagina)} issues opgehaald ({len(alle_issues)}/{totaal})")

        if len(alle_issues) >= totaal or not issues_op_pagina:
            break

        page += 1

    return alle_issues


def schrijf_naar_excel(issues: list, doelpad: Path, sonar_url: str, project_key: str):
    """Slaat de issues gestructureerd op in een Excel-bestand met openpyxl."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("Fout: openpyxl is niet geïnstalleerd. Installeer het via: pip install openpyxl")

    wb = openpyxl.Workbook()

    # ═══════════════════════════════════════════════════ BLAD 1: OVERZICHT ISSUES
    ws = wb.active
    ws.title = "Issues"
    ws.views.sheetView[0].showGridLines = True

    kolommen = [
        ("Type", 16),
        ("Ernst", 14),
        ("Status", 14),
        ("Bestand", 45),
        ("Regel", 10),
        ("Omschrijving", 60),
        ("Regelcode", 20),
        ("Inspanning", 14),
        ("Auteur", 25),
        ("Aangemaakt op", 20),
        ("Issue Sleutel", 26)
    ]

    # Header styling
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Donkerblauw
    thin_border = Border(
        left=Side(style='thin', color="E2E8F0"),
        right=Side(style='thin', color="E2E8F0"),
        top=Side(style='thin', color="E2E8F0"),
        bottom=Side(style='thin', color="E2E8F0")
    )

    # Schrijf headers
    for col_idx, (col_name, _) in enumerate(kolommen, start=1):
        cel = ws.cell(row=1, column=col_idx, value=col_name)
        cel.font = header_font
        cel.fill = header_fill
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Data rijen schrijven
    for row_idx, iss in enumerate(issues, start=2):
        comp = iss.get("component", "")
        # Verwijder projectkey prefix uit bestandspad (bijv. "Digidokters:templates/base.html" -> "templates/base.html")
        if ":" in comp:
            bestand = comp.split(":", 1)[1]
        else:
            bestand = comp

        creation_date = iss.get("creationDate", "")
        if creation_date:
            try:
                # ISO datum formatteren (2026-10-15T13:30:00+0200 -> 2026-10-15 13:30)
                dt = datetime.fromisoformat(creation_date.replace("Z", "+00:00"))
                creation_str = dt.strftime("%Y-%m-%d %H:%M")
            except Exception:
                creation_str = creation_date
        else:
            creation_str = ""

        rij_waarden = [
            iss.get("type", ""),
            iss.get("severity", ""),
            iss.get("status", ""),
            bestand,
            iss.get("line", ""),
            iss.get("message", ""),
            iss.get("rule", ""),
            iss.get("effort", ""),
            iss.get("author", ""),
            creation_str,
            iss.get("key", "")
        ]

        for col_idx, val in enumerate(rij_waarden, start=1):
            cel = ws.cell(row=row_idx, column=col_idx, value=val)
            cel.font = Font(name="Calibri", size=10)
            cel.border = thin_border
            if col_idx in (1, 2, 3, 5, 8, 10):
                cel.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cel.alignment = Alignment(horizontal="left", vertical="center")

    # Kolombreedtes instellen en autofilter
    for col_idx, (_, default_width) in enumerate(kolommen, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = default_width
    ws.auto_filter.ref = ws.dimensions

    # ═══════════════════════════════════════════════════ BLAD 2: SAMENVATTING
    ws_sum = wb.create_sheet(title="Samenvatting")
    ws_sum.views.sheetView[0].showGridLines = True

    title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    sub_font = Font(name="Calibri", size=11, bold=True)
    val_font = Font(name="Calibri", size=11)

    ws_sum.cell(row=1, column=1, value="SonarQube Issues Rapport").font = title_font
    ws_sum.cell(row=2, column=1, value=f"Project: {project_key} | Server: {sonar_url}").font = Font(italic=True, color="64748B")
    ws_sum.cell(row=3, column=1, value=f"Datum van export: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").font = Font(italic=True, color="64748B")

    # Telling per type
    types_count = {}
    severities_count = {}
    for iss in issues:
        t = iss.get("type", "ONBEKEND")
        s = iss.get("severity", "ONBEKEND")
        types_count[t] = types_count.get(t, 0) + 1
        severities_count[s] = severities_count.get(s, 0) + 1

    ws_sum.cell(row=5, column=1, value="Totaal aantal openstaande issues:").font = sub_font
    ws_sum.cell(row=5, column=2, value=len(issues)).font = Font(bold=True, size=12, color="DC2626" if len(issues) > 0 else "16A34A")

    # Tabel Type
    ws_sum.cell(row=7, column=1, value="Verdeling per Type").font = sub_font
    ws_sum.cell(row=7, column=2, value="Aantal").font = sub_font
    rij = 8
    for t_naam, count in sorted(types_count.items()):
        ws_sum.cell(row=rij, column=1, value=t_naam).font = val_font
        ws_sum.cell(row=rij, column=2, value=count).font = val_font
        rij += 1

    # Tabel Severity
    rij += 1
    ws_sum.cell(row=rij, column=1, value="Verdeling per Ernst (Severity)").font = sub_font
    ws_sum.cell(row=rij, column=2, value="Aantal").font = sub_font
    rij += 1
    for s_naam, count in sorted(severities_count.items()):
        ws_sum.cell(row=rij, column=1, value=s_naam).font = val_font
        ws_sum.cell(row=rij, column=2, value=count).font = val_font
        rij += 1

    ws_sum.column_dimensions["A"].width = 35
    ws_sum.column_dimensions["B"].width = 15

    wb.save(doelpad)
    print(f"✓ Excel-bestand succesvol opgeslagen: {doelpad}")


def ruim_oude_bestanden_op(map_pad: Path, bewaartermijn_dagen: int):
    """Verwijdert Excel-bestanden in de doelmap die ouder zijn dan de bewaartermijn."""
    if not map_pad.exists():
        return

    grensdatum = datetime.now() - timedelta(days=bewaartermijn_dagen)
    grens_timestamp = grensdatum.timestamp()

    verwijderd = 0
    for bestand in map_pad.glob("*.xlsx"):
        is_oud = False

        # 1. Controleer op datum in de bestandsnaam (bijv. sonar_issues_2026-05-10.xlsx)
        match = re.search(r'(\d{4}-\d{2}-\d{2})', bestand.name)
        if match:
            try:
                bestandsdatum = datetime.strptime(match.group(1), "%Y-%m-%d")
                if bestandsdatum < grensdatum:
                    is_oud = True
            except ValueError:
                pass

        # 2. Fallback op bestandsmodificatietijd
        if not is_oud:
            mtime = os.path.getmtime(bestand)
            if mtime < grens_timestamp:
                is_oud = True

        if is_oud:
            try:
                bestand.unlink()
                print(f"  - Verwijderd (ouder dan 3 maanden): {bestand.name}")
                verwijderd += 1
            except Exception as e:
                print(f"  ! Kon {bestand.name} niet verwijderen: {e}")

    if verwijderd == 0:
        print("Geen bestanden ouder dan 3 maanden aangetroffen.")
    else:
        print(f"✓ {verwijderd} oud(e) exportbestand(en) opgeruimd.")


def main():
    print("=" * 65)
    print(" SonarQube Issues Exporter")
    print("=" * 65)

    # 1. Lees configuratie uit .env
    env_vars = laad_env(ENV_FILE)
    sonar_token = env_vars.get("SONAR_TOKEN") or os.environ.get("SONAR_TOKEN")
    if not sonar_token:
        sys.exit("Fout: SONAR_TOKEN ontbreekt in .env en omgevingsvariabelen.")

    sonar_url = (
        env_vars.get("SONAR_HOST_URL")
        or env_vars.get("SONAR_URL")
        or os.environ.get("SONAR_HOST_URL")
        or DEFAULT_SONAR_URL
    )
    project_key = env_vars.get("SONAR_PROJECT_KEY") or DEFAULT_PROJECT_KEY

    # 2. Zorg dat de doelmap bestaat
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # 3. Haal issues op
    try:
        issues = haal_alle_issues_op(sonar_url, project_key, sonar_token)
    except Exception as e:
        sys.exit(f"Fout tijdens ophalen van issues: {e}")

    print(f"In totaal {len(issues)} openstaande issues gevonden.")

    # 4. Exporteer naar Excel
    vandaag_str = datetime.now().strftime("%Y-%m-%d")
    bestandsnaam = f"sonar_issues_{vandaag_str}.xlsx"
    doelpad = OUTPUT_DIR / bestandsnaam

    schrijf_naar_excel(issues, doelpad, sonar_url, project_key)

    # 5. Oude bestanden opschonen (ouder dan 3 maanden)
    print(f"\nOpruimen van exports ouder dan {BEWAARTERMIJN_DAGEN} dagen in '{OUTPUT_DIR.name}'...")
    ruim_oude_bestanden_op(OUTPUT_DIR, BEWAARTERMIJN_DAGEN)

    print("\nVoltooid!")


if __name__ == "__main__":
    main()
