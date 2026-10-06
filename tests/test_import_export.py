import os
from tests.base import BaseTestCase
from flask import current_app
from extensions import db
from models.organisatie import Organisatie, UserOrganisatie
from models.user import User
from werkzeug.security import generate_password_hash

class TestImportExport(BaseTestCase):
    def setUp(self):
        super().setUp()
        
        # 1. Create a second organization and admin user
        self.org2 = Organisatie(
            naam="Other Tenant",
            slug="other-tenant",
            actief=True
        )
        db.session.add(self.org2)
        db.session.commit()
        
        self.admin_user2 = User(
            naam="AdminOther",
            email="admin2@test.com",
            wachtwoord_hash=generate_password_hash("password123"),
            rol="beheerder",
            actief=True,
            moet_wachtwoord_wijzigen=False
        )
        db.session.add(self.admin_user2)
        db.session.commit()
        
        membership2 = UserOrganisatie(
            user_id=self.admin_user2.id,
            organisatie_id=self.org2.id,
            rol="beheerder",
            actief=True
        )
        db.session.add(membership2)
        db.session.commit()

    def test_import_log_access_control(self):
        # 1. Create dummy import log files for org1 (self.org.id) and org2 (self.org2.id)
        org1_log = f"import_{self.org.id}_20260729_120000.log"
        org2_log = f"import_{self.org2.id}_20260729_120000.log"
        
        org1_log_pad = os.path.join(current_app.config['IMPORT_LOG_FOLDER'], org1_log)
        org2_log_pad = os.path.join(current_app.config['IMPORT_LOG_FOLDER'], org2_log)
        
        with open(org1_log_pad, 'w') as f:
            f.write("Log contents for Org 1")
        with open(org2_log_pad, 'w') as f:
            f.write("Log contents for Org 2")
            
        try:
            # --- Scenario A: Admin of Org 1 attempts to download their own log ---
            self.login("AdminMark", "password123")
            self.select_organisatie(self.org.id)
            
            response = self.client.get(f'/importeer/log/{org1_log}')
            self.assertEqual(response.status_code, 200)
            self.assertIn("Log contents for Org 1", response.data.decode('utf-8'))
            response.close()
            
            # --- Scenario B: Admin of Org 1 attempts to download Org 2's log ---
            response = self.client.get(f'/importeer/log/{org2_log}')
            self.assertEqual(response.status_code, 403)
            response.close()
            
            self.logout()
            
            # --- Scenario C: Admin of Org 2 attempts to download Org 2's log ---
            self.login("AdminOther", "password123")
            self.select_organisatie(self.org2.id)
            
            response = self.client.get(f'/importeer/log/{org2_log}')
            self.assertEqual(response.status_code, 200)
            self.assertIn("Log contents for Org 2", response.data.decode('utf-8'))
            response.close()
            
            # --- Scenario D: Admin of Org 2 attempts to download Org 1's log ---
            response = self.client.get(f'/importeer/log/{org1_log}')
            self.assertEqual(response.status_code, 403)
            response.close()
            
        finally:
            # Clean up files
            if os.path.exists(org1_log_pad):
                os.remove(org1_log_pad)
            if os.path.exists(org2_log_pad):
                os.remove(org2_log_pad)

    def test_export_met_vraagclassificatie(self):
        """Test dat exporteer CSV en XLSX de vraagclassificatie bevatten."""
        from datetime import date
        import io
        import openpyxl
        from models.registration import Registration
        from models.digidokter import Digidokter
        from models.question_category import QuestionCategory
        from models.question_classification import QuestionClassification
        from utils.export_handler import exporteer_csv, exporteer_xlsx

        # Maak categorie, digidokter, registratie en classificatie
        cat = QuestionCategory(naam="Hardware & Toestellen", omschrijving="Vragen over hardware", volgorde=1, actief=True)
        db.session.add(cat)
        db.session.commit()

        dd = Digidokter(naam="Jan Digidokter", organisatie_id=self.org.id)
        db.session.add(dd)
        db.session.commit()

        reg = Registration(
            registratienummer="2026-9001",
            datum=date(2026, 10, 6),
            client="Test Bezoeker",
            digidokter_id=dd.id,
            leeftijdscategorie_id=self.age_category.id,
            toestel_id=self.device.id,
            herkomst_id=self.herkomst.id,
            onderwerp="Mijn laptop start niet meer op",
            organisatie_id=self.org.id,
            aangemaakt_door_id=self.admin_user.id
        )
        db.session.add(reg)
        db.session.commit()

        qc = QuestionClassification(
            registration_id=reg.id,
            category_id=cat.id,
            zekerheid=0.95,
            toelichting="Betreft een hardwareprobleem met laptop."
        )
        db.session.add(qc)
        db.session.commit()

        with self.app.test_request_context():
            from flask import session
            session['organisatie_id'] = self.org.id

            # Test CSV
            csv_bytes = exporteer_csv()
            csv_str = csv_bytes.decode('utf-8-sig')
            self.assertIn("Vraagclassificatie", csv_str)
            self.assertIn("Hardware & Toestellen", csv_str)
            self.assertIn("Mijn laptop start niet meer op", csv_str)

            # Test XLSX
            xlsx_bytes = exporteer_xlsx()
            wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
            ws = wb.active
            headers = [cell.value for cell in ws[1]]
            self.assertIn("Vraagclassificatie", headers)
            row_values = [cell.value for cell in ws[2]]
            self.assertIn("Hardware & Toestellen", row_values)

